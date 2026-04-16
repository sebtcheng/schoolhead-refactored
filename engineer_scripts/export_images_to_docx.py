"""
export_images_to_docx.py
──────────────────────────────────────────────────────────────────────────────
Exports all images from engineer_image (stored as WebP bytea in
unified_binaries) into a Word (.docx) file, grouped by project.

Each project section shows:
  - School Name, Region, Division, Municipality, Province, Project Name
  - All Internal photos, then all External photos (labelled + dated)

Output file: engineer_scripts/engineer_images_export_YYYYMMDD_HHMMSS.docx

Requirements:
    pip install psycopg2-binary python-docx Pillow

Usage:
    python engineer_scripts/export_images_to_docx.py

Optional filters (edit constants below):
    IPC_FILTER    – set to a specific IPC string to export one project only
    LIMIT_PROJECTS – set to an integer to cap the number of projects exported
──────────────────────────────────────────────────────────────────────────────
"""

import psycopg2
import os
import re
import io
from datetime import datetime

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image

# ── Config ────────────────────────────────────────────────────────────────────
IPC_FILTER     = None   # e.g. 'INF-01-2025-00048'  or None for all
LIMIT_PROJECTS = None   # e.g. 50  or None for all
IMAGE_WIDTH_IN = 3.0    # image width in inches inside the document


# ── Helpers ───────────────────────────────────────────────────────────────────

def get_db_url():
    env_path = os.path.join(os.path.dirname(__file__), '..', '.env')
    if not os.path.exists(env_path):
        raise FileNotFoundError(f".env file not found at {env_path}")
    with open(env_path, 'r') as f:
        match = re.search(r'DATABASE_URL\s*=\s*(.*)', f.read())
        if match:
            return match.group(1).strip()
    return None


def parse_db_url(url):
    pattern = r'postgres(?:ql)?://([^:]+):([^@]+)@([^:/]+):?(\d*)/(.+)'
    m = re.match(pattern, url)
    if not m:
        raise ValueError(f"Cannot parse DATABASE_URL: {url}")
    user, password, host, port, dbname = m.groups()
    return {
        'host':     host,
        'port':     int(port) if port else 5432,
        'dbname':   dbname,
        'user':     user,
        'password': password,
        'sslmode':  'require',
    }


def webp_to_png_bytes(raw_bytes):
    """Convert any image bytes (WebP, JPEG, PNG…) to PNG bytes for docx."""
    img = Image.open(io.BytesIO(bytes(raw_bytes)))
    out = io.BytesIO()
    img.convert('RGB').save(out, format='PNG')
    out.seek(0)
    return out


def add_heading(doc, text, level=1, color=None):
    p = doc.add_heading(text, level=level)
    if color:
        for run in p.runs:
            run.font.color.rgb = RGBColor(*color)
    return p


def add_meta_table(doc, meta):
    """Add a 2-column metadata table for the project details."""
    fields = [
        ('School Name',  meta.get('school_name')  or '—'),
        ('Project Name', meta.get('project_name') or '—'),
        ('Region',       meta.get('region')        or '—'),
        ('Division',     meta.get('division')      or '—'),
        ('Province',     meta.get('province')      or '—'),
        ('Municipality', meta.get('municipality')  or '—'),
        ('IPC',          meta.get('ipc')           or '—'),
        ('Category',     meta.get('category_summary') or '—'),
    ]

    table = doc.add_table(rows=len(fields), cols=2)
    table.style = 'Table Grid'

    for i, (label, value) in enumerate(fields):
        row = table.rows[i]
        # Label cell
        lc = row.cells[0]
        lc.text = label
        lc.paragraphs[0].runs[0].bold = True
        lc.paragraphs[0].runs[0].font.size = Pt(9)
        # shade label cell light blue
        tc_pr = lc._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), 'DCE6F1')
        tc_pr.append(shd)
        # Value cell
        vc = row.cells[1]
        vc.text = str(value)
        vc.paragraphs[0].runs[0].font.size = Pt(9)

    doc.add_paragraph()  # spacing after table


def add_image_row(doc, img_bytes, category, uploaded_at):
    """Add a single image with its label underneath."""
    try:
        png_stream = webp_to_png_bytes(img_bytes)
        doc.add_picture(png_stream, width=Inches(IMAGE_WIDTH_IN))
        last_para = doc.paragraphs[-1]
        last_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    except Exception as e:
        p = doc.add_paragraph(f'[Image could not be rendered: {e}]')
        p.runs[0].font.color.rgb = RGBColor(200, 0, 0)

    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = caption.add_run(
        f"{'🔵' if category == 'Internal' else '🟢'} {category or 'Uncategorized'}  |  "
        f"Uploaded: {uploaded_at.strftime('%b %d, %Y %I:%M %p') if uploaded_at else '—'}"
    )
    run.font.size = Pt(8)
    run.font.italic = True
    run.font.color.rgb = RGBColor(80, 80, 80)
    doc.add_paragraph()  # spacing


# ── Main ──────────────────────────────────────────────────────────────────────

def export_images():
    db_url = get_db_url()
    if not db_url:
        print("❌  DATABASE_URL not found in .env")
        return

    conn = psycopg2.connect(**parse_db_url(db_url))
    cur  = conn.cursor()

    # ── 1. Get distinct projects that have images ─────────────────────────────
    ipc_clause = f"AND ei.ipc = '{IPC_FILTER}'" if IPC_FILTER else ""
    limit_clause = f"LIMIT {LIMIT_PROJECTS}" if LIMIT_PROJECTS else ""

    cur.execute(f"""
        SELECT
            ei.ipc,
            ef.school_name,
            ef.region,
            ef.division,
            ef.province,
            ef.municipality,
            ef.project_name,
            COUNT(*)                                          AS total_images,
            COUNT(*) FILTER (WHERE ei.category = 'Internal') AS internal_count,
            COUNT(*) FILTER (WHERE ei.category = 'External') AS external_count
        FROM engineer_image ei
        LEFT JOIN (
            SELECT DISTINCT ON (ipc)
                ipc, school_name, region, division, province,
                municipality, project_name
            FROM engineer_form
            ORDER BY ipc, project_id DESC
        ) ef ON ef.ipc = ei.ipc
        WHERE ei.binary_id IS NOT NULL
        {ipc_clause}
        GROUP BY ei.ipc, ef.school_name, ef.region, ef.division,
                 ef.province, ef.municipality, ef.project_name
        ORDER BY ef.region, ef.division, ei.ipc
        {limit_clause}
    """)

    projects = cur.fetchall()
    proj_cols = [d[0] for d in cur.description]

    if not projects:
        print("  ✅  No images found matching your filter.")
        cur.close(); conn.close()
        return

    print(f"\n  Found {len(projects)} project(s) with images.\n")

    # ── 2. Build Word document ────────────────────────────────────────────────
    doc = Document()

    # -- Cover page
    title = doc.add_heading('Engineer Image Report', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph(
        f"Generated: {datetime.now().strftime('%B %d, %Y  %I:%M %p')}\n"
        f"Total projects: {len(projects)}"
    )
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()

    # ── 3. Per-project sections ───────────────────────────────────────────────
    for proj_idx, proj_row in enumerate(projects, 1):
        meta = dict(zip(proj_cols, proj_row))
        ipc  = meta['ipc'] or 'NO-IPC'
        meta['category_summary'] = (
            f"Internal: {meta['internal_count']}  |  "
            f"External: {meta['external_count']}  |  "
            f"Total: {meta['total_images']}"
        )

        print(f"  [{proj_idx}/{len(projects)}] {ipc}  — {meta['total_images']} image(s)")

        add_heading(doc, f"{ipc}", level=1, color=(31, 73, 125))
        add_meta_table(doc, meta)

        # Fetch images for this project — Internal first, then External
        cur.execute("""
            SELECT
                ei.category,
                ei.created_at,
                ub.content,
                ub.mime_type
            FROM engineer_image ei
            JOIN unified_binaries ub ON ub.id = ei.binary_id
            WHERE ei.ipc = %s
              AND ei.binary_id IS NOT NULL
              AND ub.content IS NOT NULL
            ORDER BY
                CASE ei.category WHEN 'Internal' THEN 1
                                  WHEN 'External' THEN 2
                                  ELSE 3 END,
                ei.created_at
        """, (ipc,))

        images = cur.fetchall()

        current_category = None
        for img_row in images:
            category, uploaded_at, raw_content, mime_type = img_row

            # Category sub-heading
            if category != current_category:
                current_category = category
                sub_label = category or 'Uncategorized'
                p = doc.add_heading(sub_label, level=2)
                for run in p.runs:
                    run.font.color.rgb = (
                        RGBColor(0, 70, 127)   if sub_label == 'Internal'
                        else RGBColor(0, 100, 0) if sub_label == 'External'
                        else RGBColor(100, 100, 100)
                    )

            add_image_row(doc, raw_content, category, uploaded_at)

        if proj_idx < len(projects):
            doc.add_page_break()

    cur.close()
    conn.close()

    # ── 4. Save ───────────────────────────────────────────────────────────────
    timestamp   = datetime.now().strftime('%Y%m%d_%H%M%S')
    out_filename = f"engineer_images_export_{timestamp}.docx"
    out_path     = os.path.join(os.path.dirname(__file__), out_filename)

    doc.save(out_path)

    print(f"\n{'='*60}")
    print(f"  ✅  Saved → {out_path}")
    print(f"  Projects exported : {len(projects)}")
    print(f"{'='*60}\n")


if __name__ == '__main__':
    export_images()
